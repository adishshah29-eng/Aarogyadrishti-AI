"""
Reads a lab report with Gemini and returns raw transcriptions only.

The file is sent inline with the request (not via the Files API, which keeps
uploads on Google's servers for 48h) and is never written to local disk.
Gemini is asked to copy numbers and units exactly as printed; conversion and
validation happen in normalize.py.
"""
import io
import json
from typing import Optional

from pydantic import BaseModel

from src.report_ingest.fields import FIELDS

DEFAULT_MODEL = "gemini-2.5-flash"
MAX_BYTES = 10 * 1024 * 1024

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
BINARY_MIMES = {"application/pdf", "image/png", "image/jpeg", "image/webp"}
TEXT_MIMES = {"text/plain", DOCX_MIME}


class ReportReadError(Exception):
    """Raised with a message that is safe and useful to show the user."""


class _Item(BaseModel):
    field: str
    value: Optional[float] = None
    text: Optional[str] = None
    unit: Optional[str] = None
    printed_name: Optional[str] = None
    qualifier: Optional[str] = None
    evidence: Optional[str] = None


class _Extraction(BaseModel):
    is_medical_report: bool
    items: list[_Item]


def build_prompt() -> str:
    lines = []
    for spec in FIELDS:
        hint = " / ".join(spec.aliases)
        if spec.categorical:
            lines.append(f'- "{spec.key}": patient sex — put "male" or "female" in "text" (printed as: {hint})')
        else:
            lines.append(f'- "{spec.key}": {spec.label} (printed as: {hint})')
    field_list = "\n".join(lines)
    return f"""You are transcribing a patient's medical/lab report into structured data.

Extract ONLY these fields, using exactly these keys in "field":
{field_list}

Rules:
1. Copy each number exactly as printed in the patient's RESULT column. Never use
   reference/normal ranges, previous results, or limits as the value.
2. Copy the unit exactly as printed (e.g. "mg/dL", "mmol/L", "umol/L", "%").
   Do NOT convert units and do NOT calculate anything.
3. If a field is not clearly present, leave it out. Never guess.
4. Blood pressure printed like "130/85" becomes two items: systolic_bp=130, diastolic_bp=85,
   unit "mmHg".
5. For "glucose", set "qualifier" to one of: "fasting", "random", "post_prandial", "unknown".
   Only fasting/FBS/(F) results are fasting. Put HbA1c under "hba1c", never under glucose.
6. For "bun": if the report shows Blood Urea Nitrogen / BUN, set qualifier "bun".
   If it shows Urea / Blood Urea / Serum Urea (not nitrogen), still use field "bun" but set
   qualifier "urea".
7. "printed_name" is the test name exactly as printed. "evidence" is the short result line
   as printed (max 80 characters).
8. Never output the patient's name, ID numbers, phone, address, or any doctor's name,
   including inside "evidence".
9. Set "is_medical_report" to false if this is not a medical or lab report.
"""


def docx_to_text(data: bytes) -> str:
    import docx
    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _friendly_api_error(exc) -> str:
    code = getattr(exc, "code", None)
    text = str(exc).lower()
    if code in (400, 401, 403) and ("api key" in text or "api_key" in text or "permission" in text):
        return "The Gemini API key was rejected. Check the key and try again."
    if code == 429:
        return "Gemini's rate limit or quota was reached. Wait a minute, or use a key with more quota."
    if code == 413 or "too large" in text:
        return "The file is too large for Gemini to read. Try a smaller scan or a single page."
    if code and code >= 500:
        return "Gemini is temporarily unavailable. Try again in a moment."
    return "Gemini couldn't read this file. You can still fill in the form manually."


def extract(data: bytes, mime: str, api_key: str, model: str = DEFAULT_MODEL) -> list:
    """Return a list of raw item dicts. Raises ReportReadError."""
    if len(data) > MAX_BYTES:
        raise ReportReadError("This file is over 10 MB. Try a smaller scan or a single page.")
    if mime not in BINARY_MIMES | TEXT_MIMES:
        raise ReportReadError("Unsupported file type. Upload a PDF, image (PNG/JPG/WEBP), DOCX or TXT.")
    if not api_key:
        raise ReportReadError("Add a Gemini API key to read reports.")

    from google import genai
    from google.genai import errors, types

    if mime == DOCX_MIME:
        try:
            body = docx_to_text(data)
        except Exception:
            raise ReportReadError("This Word file couldn't be opened. Try saving it as PDF.")
        content = types.Part.from_text(text="REPORT TEXT:\n" + body)
    elif mime == "text/plain":
        content = types.Part.from_text(text="REPORT TEXT:\n" + data.decode("utf-8", errors="replace"))
    else:
        content = types.Part.from_bytes(data=data, mime_type=mime)

    client = genai.Client(api_key=api_key)
    try:
        response = client.models.generate_content(
            model=model,
            contents=[content, build_prompt()],
            config=types.GenerateContentConfig(
                response_mime_type="application/json",
                response_schema=_Extraction,
                temperature=0,
            ),
        )
    except errors.APIError as exc:
        raise ReportReadError(_friendly_api_error(exc))
    except Exception:
        raise ReportReadError("Couldn't reach Gemini. Check your internet connection and try again.")

    return parse_response_text(getattr(response, "text", None))


def parse_response_text(text: Optional[str]) -> list:
    if not text:
        raise ReportReadError("Gemini returned no readable result for this file.")
    try:
        parsed = _Extraction.model_validate(json.loads(text))
    except Exception:
        raise ReportReadError("Gemini's answer wasn't in the expected format. Try again.")
    if not parsed.is_medical_report:
        raise ReportReadError("This doesn't look like a medical or lab report.")
    return [item.model_dump() for item in parsed.items]
